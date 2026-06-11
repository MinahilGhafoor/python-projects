from openpyxl import Workbook, load_workbook
import os
import tkinter as tk
from tkinter import messagebox

ROOT = tk.Tk()
ROOT.geometry("900x600")

FILENAME = "todo.xlsx"

# ==========================================================
# Workbook Functions
# ==========================================================
def instantiate_workbook(task):
    wb = Workbook()
    ws = wb.active
    
    ws.title = "To do List"
    ws["A1"] = "Task"
    ws["B1"] = "Status"

    return wb, ws

def save_xlsx(wb):
    wb.save(FILENAME)


def load_workBook(file):
    wb = load_workbook(file)
    ws = wb.active

    return wb, ws

# ==========================================================
# Task Functions
# ==========================================================

def addTask():

    # popup window
    top = tk.Toplevel(ROOT)
    top.title("Add Task")

    tk.Label(ROOT, text="Enter task name:").pack()

    entry = tk.Entry(top, width=40)
    entry.pack()

    def save_task():
        task = entry.get()

        if task == "":
            messagebox.showwarning("Warning", "Task cannot be empty")
            return

        if os.path.exists(FILENAME):
            wb, ws = load_workBook(FILENAME)

        else:
            wb, ws = instantiate_workbook(task)

        ws.append([task, "Pending"])
        save_xlsx(wb)
        messagebox.showinfo("Success", f"Task '{task}' added!")
        top.destroy()

    tk.Button(top, text="Save", command=save_task).pack()
    

def viewTask():
    
    if not os.path.exists(FILENAME):
        messagebox.showinfo("showinfo", f"No tasks exist.") 
        return
    
    wb, ws = load_workBook(FILENAME)
    tk.Label(ROOT, text="Your Tasks:").pack()

    scroll_bar = tk.Scrollbar(ROOT)
    scroll_bar.pack( side=tk.RIGHT,
                fill = tk.Y )

    mylist = tk.Listbox(ROOT, 
                yscrollcommand = scroll_bar.set )

            

    for idx, row in enumerate(ws.iter_rows(min_row = 2, values_only=True), start=1):
        mylist.insert(tk.END, f"{idx}- {row} ")

    mylist.pack( side = tk.LEFT, fill = tk.BOTH )
    scroll_bar.config( command = mylist.yview )

def markTask():
    
    if not os.path.exists(FILENAME):
        messagebox.showinfo("showinfo", f"No tasks exist.")
        return
    
    wb, ws = load_workBook(FILENAME)
    
    var = IntVar()

    for idx, row in enumerate(ws.iter_rows(min_row = 2, values_only=True), start=1):
        chk = tk.Checkbutton(ROOT, text=row, variable=var).pack()

    try:
        var 

    except (ValueError, IndexError):
            print("Invalid choice.")


def delTask():
    pass
    """if not os.path.exists(FILENAME):
        print("No tasks yet!")
        return
    
    wb, ws = load_workBook(FILENAME)
    print("\nYour Tasks.")

    for idx, row in enumerate(ws.iter_rows(min_row = 2, values_only=True), start=1):
        print(idx,row)

    try:
            choice = int(input("Enter the task number to delete: "))

            row_num = choice + 1

            ws.delete_rows(row_num)

            wb.save(FILENAME)

    except (ValueError, IndexError):
        print("Invalid choice.")"""

def menu():

    values = {"Add a new Task" : addTask ,
            "View all tasks" : viewTask,
            "Mark task as complete" : markTask,
            "Delete a task" : delTask,
            "Quit" : ROOT.destroy}

    for text, cmd in values.items():
        tk.Button(ROOT, text=text, command=cmd).pack(pady=10)

    
menu()

ROOT.mainloop() 
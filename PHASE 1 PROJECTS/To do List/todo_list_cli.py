from openpyxl import Workbook, load_workbook
import os


FILENAME = "todo.xlsx"

def instantiate_workbook(task):
    wb = Workbook()
    ws = wb.active
    
    ws.title = "To do List.."
    ws["A1"] = "Task"
    ws["B1"] = "Status"
    return wb, ws

def save_xlsx(wb, filename):
    filename = f"{filename}.xlsx"
    wb.save(filename)

    return filename

def load_workBook(file):
    wb2 = load_workbook(file)
    ws2 = wb2.active

    return wb2, ws2

def read_workBook(ws2):
    print("\nReading data:")
    
    for row in ws2.iter_rows(values_only=True):
        print(row)
        

def addTask():
    task = input("Enter the name of task you want in your list:  ")

    if os.path.exists(FILENAME):
        wb, ws = load_workBook(FILENAME)
    
    else:
        wb, ws = instantiate_workbook(task)

    ws.append([task, "Pending"])
    save_xlsx(wb, "todo")
    print(f"Task '{task}' added!")
    

def viewTask():
    
    if not os.path.exists(FILENAME):
        print("No tasks yet!")
        return
    
    wb, ws = load_workBook(FILENAME)
    print("\nYour Tasks:")

    for idx, row in enumerate(ws.iter_rows(min_row = 2, values_only=True), start=1):
        print(idx,row)

def markTask():
    if not os.path.exists(FILENAME):
        print("No tasks yet!")
        return
    
    wb, ws = load_workBook(FILENAME)
    print("\nYour Tasks.")

    for idx, row in enumerate(ws.iter_rows(min_row = 2, values_only=True), start=1):
        print(idx,row)

    try:
            choice = int(input("Enter the task number to mark complete: "))

            row_num = choice + 1

            ws.cell(row=row_num, column=2, value="Complete")
            print("Marked!")
            wb.save(FILENAME)

    except (ValueError, IndexError):
            print("Invalid choice.")


def delTask():
    if not os.path.exists(FILENAME):
        print("No tasks yet!")
        return
    
    wb, ws = load_workBook(FILENAME)
    print("\nYour Tasks.")

    for idx, row in enumerate(ws.iter_rows(min_row = 2, values_only=True), start=1):
        print(idx,row)

    try:
            choice = int(input("Enter the task number to delete: "))

            row_num = choice + 1

            ws.delete_rows(row_num)

            wb.save(FILENAME)

    except (ValueError, IndexError):
        print("Invalid choice.")

def menu():
    while True:

        choice = input(f"Want to:\n1. Add Tasks.\n2. View all tasks.\n3. Mark task as complete.\n4. Delete a task.\n5. quit (1,2,3,4,5) ::::: ")

        if choice == "1":
            addTask()

        elif choice == "2":
            viewTask()

        elif choice == "3":
            markTask()

        elif choice == "4":
            delTask()

        elif choice == "5":
            print("Bye!")
            break

        else:
            print("Invalid Input!. Try AGain")

menu()